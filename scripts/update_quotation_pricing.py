"""Revise FCIAZ quotation pricing.

- Drop the build invoice from K7,000 to K2,500 (further founder-friendly rate)
- Restructure TABLE 2 (pricing breakdown) to show what was actually delivered:
  6 originally-scoped line items at founder-friendly rates + 7 post-scope
  additions at K0 with their commercial reference values visible
- Rewrite the "Pricing note" paragraph to explain the founder-friendly nature
- Update Section 3 total to K2,500
- Update TABLE 5 (payment terms) amounts to K1,250 / K1,250
- Clarify in Section 4.1 that FCIAZ will register and maintain the
  infrastructure accounts directly; Manova Labs assists at no charge
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

DOC_PATH = Path(r"C:\Users\RICHARD_TEMBO\Desktop\Projects\fciaz\docs\Manova-Labs-Quotation-FCIAZ-Website-2026.docx")


def find_para_starting_with(doc: Document, prefix: str):
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    return None


def set_para_text(p, text: str) -> None:
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def replace_in_para(p, old: str, new: str) -> None:
    full = "".join(r.text for r in p.runs)
    if old not in full:
        return
    new_text = full.replace(old, new)
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)


def replace_table_in_place(doc: Document, table_index: int, new_rows: list[tuple]) -> None:
    """Delete the table at table_index and insert a fresh one at the same XML position.

    new_rows is a list of tuples; the first tuple is the header row.
    """
    old_table = doc.tables[table_index]
    old_elem = old_table._tbl
    parent = old_elem.getparent()
    prev_elem = old_elem.getprevious()
    parent.remove(old_elem)

    # Build the replacement table (appended to end first, then moved).
    n_cols = len(new_rows[0])
    new_table = doc.add_table(rows=1, cols=n_cols)
    new_table.style = "Light Grid Accent 1"
    for i, val in enumerate(new_rows[0]):
        new_table.rows[0].cells[i].text = val
    for row_data in new_rows[1:]:
        row_cells = new_table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val

    new_elem = new_table._tbl
    new_elem.getparent().remove(new_elem)
    if prev_elem is not None:
        prev_elem.addnext(new_elem)
    else:
        parent.insert(0, new_elem)


def main() -> None:
    doc = Document(DOC_PATH)

    # ---- 1. Rewrite the "Pricing note" paragraph (P04) ----
    p = find_para_starting_with(doc, "Pricing note: Manova Labs is providing")
    if p:
        set_para_text(
            p,
            "Pricing note: Manova Labs is providing this engagement at a deep founder-friendly rate "
            "in support of FCIAZ's public-health mission. The build invoice is K2,500. During the engagement, "
            "seven additional features were added beyond the original scope at no additional charge: "
            "Sanity CMS integration, dynamic news and events pages, Cloudflare + OpenNext deployment migration, "
            "the modern UI design system, full multilingual translations across three languages, "
            "security hardening, and PWA / favicon setup. The commercial reference value of all delivered work "
            "is approximately K25,400, so the founder-friendly reduction across the engagement is around 90%. "
            "The line items below show what was charged for each component alongside its commercial reference.",
        )

    # ---- 2. Update Section 3 total paragraph (P52) ----
    p = find_para_starting_with(doc, "Total: seven thousand Zambian Kwacha")
    if p:
        set_para_text(
            p,
            "Total: two thousand five hundred Zambian Kwacha (K2,500). This fee covers the developer's time "
            "and effort only and excludes VAT, hosting, and third-party service costs, which FCIAZ will "
            "register and maintain directly (see Section 4 for the annual cost breakdown).",
        )

    # ---- 3. Replace TABLE 2 (pricing breakdown) with the transparent version ----
    pricing_rows = [
        # Header
        ("Item", "Description", "Quoted (K)", "Commercial reference (K)"),
        # Originally scoped items at founder-friendly rates
        (
            "Discovery, architecture, and responsive design",
            "Requirements workshop, sitemap, page specifications, component design, brand palette, accessibility baseline.",
            "500",
            "2,500",
        ),
        (
            "Frontend build (11 pages, 3 locales, responsive)",
            "All public pages, header, footer, language switcher, mobile menu, sticky accessible navigation.",
            "1,000",
            "8,000",
        ),
        (
            "Donations page — bank + MTN MoMo + Airtel Money + Zamtel Money",
            "Manual donations layout, per-provider step-by-step instructions, pending-bank-detail notice.",
            "200",
            "1,000",
        ),
        (
            "Forms and backend (newsletter, contact, volunteer)",
            "Self-hosted Supabase storage, Resend transactional email, server actions with Zod validation, WhatsApp floating button.",
            "300",
            "1,500",
        ),
        (
            "SEO, accessibility, performance, deployment",
            "Per-locale metadata, OG image, robots.txt, semantic HTML, static pre-rendering, Cloudflare Workers setup.",
            "300",
            "1,200",
        ),
        (
            "Documentation, handoff, and 30-day post-launch support",
            "README, .env.example, Supabase schema, hand-off notes, plus 30 days of bug fixes and minor copy edits.",
            "200",
            "800",
        ),
        (
            "Subtotal — originally scoped items",
            "",
            "2,500",
            "15,000",
        ),
        # Post-scope additions, all quoted at K0
        (
            "Sanity CMS integration (added during development)",
            "Studio with i18n schemas (post, event, resource, siteSettings), GROQ queries, image pipeline, hardcoded fallback for offline operation.",
            "0",
            "2,500",
        ),
        (
            "Dynamic news and events pages (added)",
            "/news/[slug] and /events/[slug] dynamic routes backed by Sanity content.",
            "0",
            "800",
        ),
        (
            "Cloudflare + OpenNext deployment migration (added)",
            "Migrated from Vercel to Cloudflare Workers via @opennextjs/cloudflare; wrangler + open-next.config.ts; production security headers.",
            "0",
            "1,500",
        ),
        (
            "Modern UI design system (added)",
            "Mesh gradients, glass cards, aurora animations, hover micro-interactions across every page.",
            "0",
            "2,000",
        ),
        (
            "Multilingual translations across three languages (added)",
            "Full Bemba and Nyanja translation coverage, clinical phrasing flagged for FCIAZ review.",
            "0",
            "1,800",
        ),
        (
            "Security hardening (added)",
            "CSP, HSTS, Permissions-Policy, IP-based rate limiting on forms, sanitized error logs.",
            "0",
            "1,200",
        ),
        (
            "PWA and favicon setup (added)",
            "Real favicons in multiple sizes, web app manifest, installable PWA.",
            "0",
            "600",
        ),
        (
            "Subtotal — features added beyond original scope",
            "At no additional charge under the founder-friendly rate.",
            "0",
            "10,400",
        ),
        # Grand total
        (
            "TOTAL",
            "Total amount invoiced to FCIAZ",
            "K 2,500",
            "K 25,400",
        ),
    ]
    replace_table_in_place(doc, 2, pricing_rows)

    # ---- 4. Update TABLE 5 (payment terms) to reflect K2,500 total ----
    payment_table = doc.tables[5]
    if len(payment_table.rows) >= 3 and payment_table.rows[1].cells[1].text.strip() == "3,500":
        payment_table.rows[1].cells[1].text = "1,250"
        payment_table.rows[2].cells[1].text = "1,250"
    else:
        # Fallback: search and replace
        for row in payment_table.rows:
            for cell in row.cells:
                if cell.text.strip() == "3,500":
                    cell.text = "1,250"

    # ---- 5. Update the Payment-terms intro paragraph (P91) ----
    p = find_para_starting_with(doc, "All invoices are payable within seven")
    if p:
        replace_in_para(
            p,
            "Bank details will be provided on the invoice header.",
            "Bank details will be provided on the invoice header. The build invoice of K2,500 was issued at launch on 25 September 2026; future maintenance packages under Section 4 are invoiced annually in advance.",
        )

    # ---- 6. Clarify Section 4.1 that FCIAZ handles infrastructure directly ----
    p = find_para_starting_with(
        doc,
        "The following services are billed directly by their providers, not by Manova Labs.",
    )
    if p:
        set_para_text(
            p,
            "FCIAZ will register and maintain these third-party accounts directly. "
            "Manova Labs will assist with the initial setup of each one (domain registration at ZAMICA, "
            "Cloudflare Workers project, Sanity project, Supabase project, Resend account) at no additional charge. "
            "All items below are within the free tier of their respective providers as of the date of this quotation. "
            "Manova Labs will flag any move to a paid tier and obtain FCIAZ written approval before costs are incurred.",
        )

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")


if __name__ == "__main__":
    main()
